"""DOCX document extractor using python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Any

import docx

from tidyos.indexing.extractors.base import (
    BaseExtractor,
    ExtractedContent,
    normalize_extracted_text,
)
from tidyos.logging_config import get_logger

logger = get_logger("indexing.extractors.docx")

SUPPORTED_DOCX_EXTENSIONS = {".docx"}


class DocxExtractor(BaseExtractor):
    """DOCX text extractor using python-docx."""

    def extract(self, file_path: Path, max_chars: int = 12000) -> ExtractedContent:
        path_str = str(file_path.resolve())
        if not file_path.exists() or not file_path.is_file():
            return ExtractedContent(
                file_path=path_str,
                is_valid=False,
                error=f"File not found: {path_str}",
            )

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            logger.warning(f"Could not open DOCX file {file_path}: {e}")
            return ExtractedContent(
                file_path=path_str,
                is_valid=False,
                error=f"Corrupt or unreadable DOCX: {e}",
            )

        try:
            metadata: Dict[str, Any] = {}
            if hasattr(doc, "core_properties") and doc.core_properties:
                props = doc.core_properties
                metadata["title"] = props.title or None
                metadata["author"] = props.author or None
                metadata["subject"] = props.subject or None

            parts = []
            char_count = 0
            is_truncated = False

            # Extract paragraph text
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    parts.append(txt)
                    char_count += len(txt) + 1
                    if char_count >= max_chars:
                        is_truncated = True
                        break

            # Extract table cells if bounded limit allows
            if not is_truncated and hasattr(doc, "tables"):
                for tbl in doc.tables:
                    for row in tbl.rows:
                        row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if row_txt:
                            parts.append(row_txt)
                            char_count += len(row_txt) + 1
                            if char_count >= max_chars:
                                is_truncated = True
                                break
                    if is_truncated:
                        break

            combined = "\n".join(parts)
            if len(combined) > max_chars:
                combined = combined[:max_chars]
                is_truncated = True

            normalized = normalize_extracted_text(combined)

            return ExtractedContent(
                file_path=path_str,
                text=normalized,
                char_count=len(normalized),
                is_truncated=is_truncated,
                is_valid=True,
                metadata=metadata,
            )
        except Exception as e:
            logger.error(f"Error reading DOCX content from {file_path}: {e}", exc_info=True)
            return ExtractedContent(
                file_path=path_str,
                is_valid=False,
                error=f"DOCX extraction error: {e}",
            )
