"""Tests for document and content extraction pipeline."""

from pathlib import Path
import pytest
import pymupdf
import docx
from PIL import Image

from tidyos.indexing.extractors import (
    ContentExtractor,
    DocxExtractor,
    ImageExtractor,
    PdfExtractor,
    TextExtractor,
    normalize_extracted_text,
)


def test_normalize_extracted_text():
    raw = "Hello\x00\x07 World!\r\n\r\n\r\nThis   is   a   test.\n\n\n\nLine."
    cleaned = normalize_extracted_text(raw)
    assert "\x00" not in cleaned
    assert "\x07" not in cleaned
    assert "Hello World!" in cleaned
    assert "This is a test." in cleaned
    assert "\n\n\n" not in cleaned


def test_text_extractor(tmp_path: Path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Line 1\nLine 2\nLine 3", encoding="utf-8")

    extractor = TextExtractor()
    res = extractor.extract(txt_file)

    assert res.is_valid is True
    assert res.is_truncated is False
    assert "Line 1" in res.text
    assert "Line 3" in res.text
    assert res.char_count > 0


def test_text_extractor_truncation(tmp_path: Path):
    txt_file = tmp_path / "long.txt"
    txt_file.write_text("A" * 500, encoding="utf-8")

    extractor = TextExtractor()
    res = extractor.extract(txt_file, max_chars=100)

    assert res.is_valid is True
    assert res.is_truncated is True
    assert len(res.text) <= 100


def test_pdf_extractor(tmp_path: Path):
    pdf_file = tmp_path / "test.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Quarterly Financial Invoice #10293")
    doc.set_metadata({"title": "Q3 Invoice", "author": "Acme Corp"})
    doc.save(str(pdf_file))
    doc.close()

    extractor = PdfExtractor(max_pages=5)
    res = extractor.extract(pdf_file)

    assert res.is_valid is True
    assert "Quarterly Financial Invoice" in res.text
    assert res.metadata.get("title") == "Q3 Invoice"
    assert res.metadata.get("page_count") == 1
    assert res.is_truncated is False


def test_docx_extractor(tmp_path: Path):
    docx_file = tmp_path / "test.docx"
    doc = docx.Document()
    doc.core_properties.title = "Project Roadmap"
    doc.add_paragraph("TidyOS Autonomous File System Agent")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Phase 3"
    table.cell(1, 1).text = "Active"
    doc.save(str(docx_file))

    extractor = DocxExtractor()
    res = extractor.extract(docx_file)

    assert res.is_valid is True
    assert "TidyOS Autonomous File System Agent" in res.text
    assert "Milestone | Status" in res.text or "Phase 3 | Active" in res.text
    assert res.metadata.get("title") == "Project Roadmap"


def test_image_extractor(tmp_path: Path):
    img_file = tmp_path / "test.png"
    img = Image.new("RGB", (300, 200), color="blue")
    img.save(img_file, format="PNG")

    extractor = ImageExtractor()
    res = extractor.extract(img_file)

    assert res.is_valid is True
    assert res.metadata["width"] == 300
    assert res.metadata["height"] == 200
    assert "base64" in res.metadata
    assert len(res.metadata["base64"]) > 0
    assert res.metadata["format"] == "PNG"


def test_composite_content_extractor(tmp_path: Path):
    composite = ContentExtractor()

    assert composite.can_extract("doc.pdf") is True
    assert composite.can_extract("file.docx") is True
    assert composite.can_extract("notes.md") is True
    assert composite.can_extract("image.png") is True
    assert composite.can_extract("unknown.exe") is False

    # Non-existent file
    res = composite.extract(tmp_path / "non_existent.pdf")
    assert res.is_valid is False
    assert "File not found" in (res.error or "")

    # Unsupported format
    unsupported = tmp_path / "test.bin"
    unsupported.write_bytes(b"\x00\x01\x02")
    res2 = composite.extract(unsupported)
    assert res2.is_valid is False
    assert "Unsupported file format" in (res2.error or "")
